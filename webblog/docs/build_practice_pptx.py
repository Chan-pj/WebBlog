# -*- coding: utf-8 -*-
"""Generate 5-week practice PowerPoint for webblog (eGovFrame) course."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

OUT_PATH = r"d:\eGOV4.3.1\workspace-egov\webblog\docs\webblog-5week-practice.pptx"


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle and slide.placeholders[1]:
        slide.placeholders[1].text = subtitle


def add_bullet_slide(prs, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        if i == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(18)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "webblog 실습 교안 (5주)",
        "전자정부프레임워크 4.3.1 · Spring MVC · MyBatis · MariaDB · eGovFrame Web Project",
    )

    # --- 1주 ---
    add_bullet_slide(
        prs,
        "1주차 · 학습 목표",
        [
            "프로젝트 전체 흐름과 계층 구조 이해",
            "개발 환경 준비 및 eGovFrame 웹 프로젝트 구성 파악",
            "설정 파일(pom, web, Spring, MyBatis, MVC, Tiles) 역할 이해",
            "MariaDB 스키마·테이블·권한과 애플리케이션 연결",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · 프로젝트 전체 흐름",
        [
            "브라우저 요청 → Tomcat → web.xml(필터, 리스너, DispatcherServlet)",
            "DispatcherServlet → dispatcher-servlet.xml(컨트롤러 매핑, 뷰 리졸버, 인터셉터)",
            "@Controller → Service(@Service) → Mapper(MyBatis) → MariaDB",
            "응답: Tiles 정의 → layout.jsp + body JSP → HTML",
            "인증: 세션(loginUser), 관리자는 /admin/** 인터셉터로 ADMIN 역할 검사",
        ],
        "학생에게 한 번 URL을 따라가며 로그를 보라고 안내하면 좋습니다.",
    )

    add_bullet_slide(
        prs,
        "1주차 · 프로젝트 준비사항",
        [
            "JDK 8, Maven 3.x, Apache Tomcat 9.x(또는 수업 환경에 맞는 버전)",
            "MariaDB 설치 및 클라이언트(mysql, HeidiSQL 등)",
            "IDE: Eclipse + eGovFrame 플러그인 또는 VS Code + Maven(선택)",
            "Git(선택), 카카오 개발자 앱·REST API 키·Redirect URI 등록",
            "환경변수: KAKAO_REST_API_KEY (카카오 로그인 실습 시)",
            "DB 스크립트: src/main/resources/db/mariadb/*.sql 실행",
            "DB 접속 정보: database.properties 또는 ~/.webblog/database.properties",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · eGovFrame Web Project 구성",
        [
            "패키징: pom.xml에서 <packaging>war</packaging>",
            "src/main/java : 컨트롤러·서비스·매퍼 인터페이스·VO",
            "src/main/resources : Spring XML, MyBatis XML, message, db 스크립트",
            "src/main/webapp : JSP, 정적 리소스, WEB-INF/web.xml",
            "WEB-INF/config/egovframework/springmvc : MVC·Tiles 설정",
            "Root Context: web.xml의 ContextLoaderListener + context-*.xml",
            "Servlet Context: DispatcherServlet + dispatcher-servlet.xml",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · 데이터베이스 및 테이블 구성",
        [
            "DB 이름 예: pub_blog (utf8mb4)",
            "00_create_database_and_user.sql: DB·사용자·GRANT",
            "01_create_tables.sql: blog_user, blog, blog_member, post, blacklist_word",
            "02_seed_admin_optional.sql(선택): role=ADMIN 테스트 계정",
            "앱 연결: jdbc:mariadb://호스트:3306/pub_blog?...",
            "MyBatis 매퍼 XML과 컬럼명·테이블명이 일치해야 함",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · pom.xml 기능",
        [
            "의존성: eGovFrame RTE, Spring, MyBatis, MariaDB 드라이버, Tiles, JSTL 등",
            "commons-dbcp2: DataSource(BasicDataSource)용",
            "빌드: maven-compiler-plugin (Java 8), maven-war-plugin",
            "mvn package → target/webblog-1.0.0.war 생성",
            "인코딩: project.build.sourceEncoding=UTF-8",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · web.xml 기능",
        [
            "CharacterEncodingFilter: *.do 요청 UTF-8",
            "HTMLTagFilter: eGovFrame 제공 필터",
            "contextConfigLocation: classpath*:egovframework/spring/context-*.xml",
            "ContextLoaderListener: Root ApplicationContext 로딩",
            "DispatcherServlet(url-pattern /): Spring MVC 프론트 컨트롤러",
            "welcome-file: index.jsp → / 로 포워드 등",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · properties · Spring context · sql-mapper-config.xml",
        [
            "egovframework/properties/database.properties: db.url, username, password",
            "context-properties.xml: PropertySourcesPlaceholderConfigurer로 치환",
            "context-datasource.xml: dataSource 빈 → DBCP2 BasicDataSource",
            "context-common.xml: component-scan(서비스 등), messageSource",
            "context-mapper.xml: SqlSessionFactory, Mapper 스캔",
            "sql-mapper-config.xml: MyBatis typeAlias 또는 설정(프로젝트 기준)",
            "mapper XML: egovframework/sqlmap/mappers/*.xml",
        ],
    )

    add_bullet_slide(
        prs,
        "1주차 · dispatcher-servlet.xml · tiles.xml",
        [
            "dispatcher-servlet.xml: @Controller 스캔(egovframework, com.chan)",
            "TilesConfigurer + TilesViewResolver: home, auth/*, admin/*, blog/*, post/*",
            "UrlBasedViewResolver(순서 2): 샘플 egovframework/example JSP 경로",
            "LocaleChangeInterceptor, 예외 SimpleMappingExceptionResolver",
            "MappedInterceptor: /admin/** → AdminAuthInterceptor",
            "tiles.xml: defaultLayout → header, menu, footer, body 속성",
        ],
    )

    # --- 2주 ---
    add_bullet_slide(
        prs,
        "2주차 · 학습 목표",
        [
            "OAuth 2.0 인가 코드 흐름 이해",
            "카카오 로그인: state, 토큰, 사용자 정보, 세션 연동",
            "웹 계층별 역할: Controller / Service / Mapper / XML / JSP",
        ],
    )

    add_bullet_slide(
        prs,
        "2주차 · 카카오 인증 로그인 구현 구조",
        [
            "시작: GET /auth/kakao/login.do → state 세션 저장 → 카카오 인가 URL redirect",
            "콜백: GET /auth/kakao/callback.do?code&state",
            "state 검증 후 POST https://kauth.kakao.com/oauth/token (access_token)",
            "GET https://kapi.kakao.com/v2/user/me (Bearer)",
            "kakao id, email, nickname 추출 → registerKakaoUser (없으면 insert)",
            "session.setAttribute(\"loginUser\", user) → redirect /",
            "환경변수 KAKAO_REST_API_KEY, Redirect URI 콘솔 등록 필수",
        ],
    )

    add_bullet_slide(
        prs,
        "2주차 · 컨트롤러 · AuthController",
        [
            "패키지: com.chan.web.AuthController",
            "일반: GET/POST /auth/login.do, /auth/register.do, /auth/logout.do",
            "카카오: /auth/kakao/login.do, /auth/kakao/callback.do",
            "RestTemplate + Jackson으로 토큰·사용자 JSON 처리",
            "뷰 이름: auth/login, auth/register (Tiles)",
        ],
    )

    add_bullet_slide(
        prs,
        "2주차 · 서비스 · 매퍼 · XML · JSP",
        [
            "UserService / UserServiceImpl: registerUser, registerKakaoUser, getUserByEmail",
            "UserMapper.java + UserMapper.xml: blog_user CRUD·조회",
            "registerKakaoUser: kakao_id 중복 시 기존 사용자 반환",
            "JSP: WEB-INF/jsp/auth/login.jsp, register.jsp",
            "menu.jsp: sessionScope.loginUser 유무로 로그인/로그아웃 표시",
        ],
    )

    # --- 3주 ---
    add_bullet_slide(
        prs,
        "3주차 · 학습 목표",
        [
            "블로그·멤bers·설정 변경 흐름",
            "BlogService와 BlogMapper, 관련 JSP 연결",
            "관리자 영역과 일반 블로그 화면 구분",
        ],
    )

    add_bullet_slide(
        prs,
        "3주차 · 블로그 생성 및 관리 구조",
        [
            "회원 가입 시 UserServiceImpl.registerUser → createDefaultBlog",
            "blog 테이블 insert + 기본 한도(게시판 수, 파일 용량 등)",
            "BlogController: /blog/{blogId}/view.do, members.do, member/status.do",
            "AdminController: /admin/dashboard.do, blogs.do, blog/update.do",
            "AdminAuthInterceptor: 세션 + role==ADMIN",
        ],
    )

    add_bullet_slide(
        prs,
        "3주차 · 계층별 파일 (블로그·관리자)",
        [
            "Controller: BlogController.java, AdminController.java",
            "Service: BlogService, BlogServiceImpl",
            "Mapper: BlogMapper.java + BlogMapper.xml",
            "JSP: blog/blogView.jsp, blog/memberList.jsp",
            "JSP: admin/dashboard.jsp, admin/blogList.jsp",
            "게시글: PostController, PostService, PostMapper, post/*.jsp (연계)",
        ],
    )

    # --- 4주 ---
    add_bullet_slide(
        prs,
        "4주차 · 학습 목표",
        [
            "JUnit 5 기반 단위·통합 테스트 개념",
            "Spring TestContext로 서비스·매퍼 테스트",
            "인증·권한·예외 시나리오를 테스트 케이스로 정리",
        ],
    )

    add_bullet_slide(
        prs,
        "4주차 · 테스트 케이스 구현 (권장 범위)",
        [
            "도구: JUnit Jupiter, spring-test (pom.xml에 test scope)",
            "UserServiceImpl: 이메일 중복 가입, 카카오 id 중복 시 기존 사용자",
            "PostServiceImpl: 금지어·개인정보 검증 시 예외(ModerationService mock)",
            "Mapper 통합 테스트: @SpringJUnitConfig + 실제 또는 테스트 DB(H2·Testcontainers 선택)",
            "실행: mvn test (또는 IDE에서 테스트 클래스 실행)",
            "과제: 관리자 인터셉터 경로에 대한 MockMvc 테스트 추가(심화)",
        ],
    )

    add_bullet_slide(
        prs,
        "4주차 · 실습 과제 예시",
        [
            "Given-When-Then 형식으로 시나리오 3개 이상 작성",
            "회원가입 후 블로그 기본 생성 여부 검증",
            "카카오 사용자 정보 JSON 파싱 로직 단위 테스트(순수 Java)",
            "테스트 DB 분리: application-test.properties 또는 전용 스키마",
        ],
    )

    # --- 5주 ---
    add_bullet_slide(
        prs,
        "5주차 · 학습 목표",
        [
            "2-tier: 애플리케이션 서버(WAS)와 DB 서버 분리",
            "운영 배포: 패키징, 전송, 설정, 보안(방화벽·계정)",
            "eGovFrame WAR 배포 및 장애 대응 포인트",
        ],
    )

    add_bullet_slide(
        prs,
        "5주차 · 2-tier 구현 구조",
        [
            "Tier 1: Tomcat 등 WAS — webblog-*.war 배포",
            "Tier 2: MariaDB — 별도 호스트, 포트 3306(또는 변경)",
            "앱은 jdbc:mariadb://DB서버IP:3306/pub_blog 로 접속",
            "database.properties를 서버 환경에 맞게 배치(~/.webblog/ 등)",
            "카카오 Redirect URI를 운영 도메인으로 변경",
            "방화벽: WAS → DB 포트만 허용, DB는 외부 직접 노출 최소화",
        ],
    )

    add_bullet_slide(
        prs,
        "5주차 · 서버(WAS)에 설치·설정할 항목",
        [
            "JRE/JDK 8 이상 (Tomcat 실행용)",
            "Apache Tomcat: webapps에 WAR 배포, context path 확인",
            "환경변수: KAKAO_REST_API_KEY 등",
            "로그: logs/catalina.out 등으로 기동 오류 확인",
            "타임존·UTF-8, 파일 업로드 경로(향후 확장 시) 검토",
        ],
    )

    add_bullet_slide(
        prs,
        "5주차 · DB 서버 설치·설정",
        [
            "MariaDB 서버 패키지 설치 및 서비스 기동",
            "00·01 SQL로 DB·계정·테이블 생성",
            "bind-address, 방화벽, 최소 권한 계정",
            "문자셋 utf8mb4, 백업 정책(선택)",
        ],
    )

    add_bullet_slide(
        prs,
        "5주차 · 컴파일 · 전달 · 실행",
        [
            "개발 PC: mvn clean package -DskipTests (또는 테스트 포함)",
            "산출물: target/webblog-1.0.0.war",
            "전달: SCP, SFTP, CI 파이프라인, 사내 배포 도구 등",
            "서버: $CATALINA_HOME/webapps/ 에 복사 후 기동 또는 manager 배포",
            "기동 후 http://서버:8080/컨텍스트/ 접속 확인",
            "장애 시: ClassNotFoundException → 의존성 누락, DB 연결 실패 → 방화벽·계정·URL",
        ],
    )

    add_bullet_slide(
        prs,
        "마무리 · 참고 자료 (프로젝트 내)",
        [
            "README.md: 전체 구성, URL, DB·카카오 설정",
            "docs/webblog-5week-practice.pptx: 본 교안",
            "db/mariadb/*.sql: DB 초기화",
            "공식 문서: eGovFrame, Spring MVC, MyBatis, Kakao Developers",
        ],
    )

    prs.save(OUT_PATH)
    print("Saved:", OUT_PATH)


if __name__ == "__main__":
    main()
