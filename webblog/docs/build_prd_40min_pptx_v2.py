# -*- coding: utf-8 -*-
"""Generate styled 40-minute / 20-page PRD lecture slides (v2)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT_PATH = r"d:\eGov4.3.1\workspace-egov\webblog\docs\수업발표용_40분_20페이지_블로그PRD_v2.pptx"

COLOR_PRIMARY = RGBColor(28, 63, 170)
COLOR_ACCENT = RGBColor(34, 139, 230)
COLOR_TEXT = RGBColor(33, 37, 41)
COLOR_LIGHT_BG = RGBColor(245, 248, 255)


def style_title(slide, title_text, subtitle_text=""):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.8), Inches(1.2))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.8), Inches(1.0))
        stf = sub_box.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle_text
        sp.font.size = Pt(20)
        sp.font.color.rgb = COLOR_TEXT


def add_footer(slide, page_no):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_LIGHT_BG
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(11.6), Inches(7.18), Inches(1.4), Inches(0.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = f"{page_no}/20"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY


def add_section_slide(prs, page_no, section_title, section_subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_LIGHT_BG
    bg.line.fill.background()
    style_title(slide, section_title, section_subtitle)
    add_footer(slide, page_no)


def add_content_slide(prs, page_no, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    # title bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.16), Inches(12), Inches(0.6))
    ttf = tbox.text_frame
    ttf.clear()
    tp = ttf.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(255, 255, 255)

    # bullet area
    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.25), Inches(11.9), Inches(5.6))
    btf = body.text_frame
    btf.clear()
    for i, line in enumerate(bullets):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        p.text = f"- {line}"
        p.font.size = Pt(22)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(12)

    # accent line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.22), Inches(0.12), Inches(5.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT
    accent.line.fill.background()

    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    add_footer(slide, page_no)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1
    add_section_slide(prs, 1, "PRD 기반 블로그 플랫폼 설계와 구현", "수업 발표 자료 (40분 · 20페이지)")

    slides = [
        ("왜 PRD가 필요한가", ["요구사항 누락과 해석 차이를 줄인다", "팀의 공통 언어를 만든다", "테스트 기준(수용 기준)을 명확히 한다"]),
        ("PRD의 핵심 구성요소", ["문제 정의와 목표", "범위(In Scope / Out of Scope)", "사용자 역할/시나리오", "기능·비기능 요구사항", "수용 기준과 성공 지표"]),
        ("프로젝트 배경과 문제", ["권한 분리가 약하면 운영 안정성이 떨어진다", "승인 절차가 없으면 관리자 권한 오남용 가능", "업로드 정책이 없으면 저장소/비용 통제가 어렵다"]),
        ("이해관계자와 사용자", ["슈퍼관리자: 승인/정책 통제", "블로그관리자: 블로그 운영", "블로그멤버: 열람·작성", "강사/조교: 학습 성과 검증"]),
        ("역할 모델(RBAC) 설계", ["SUPER_ADMIN", "BLOG_ADMIN", "BLOG_MEMBER", "승인 상태: is_approved"]),
        ("인증/가입 요구사항", ["이메일 로그인", "카카오 OAuth 로그인", "가입 유형 선택(BLOG_ADMIN / BLOG_MEMBER)", "BLOG_ADMIN은 승인 전 대기"]),
        ("로그인 후 라우팅 정책", ["SUPER_ADMIN -> 관리자 대시보드", "승인 전 BLOG_ADMIN -> 대기 화면", "BLOG_MEMBER -> 소속 블로그"]),
        ("블로그멤버 글쓰기 정책", ["소유자 또는 ACTIVE 멤버만 글쓰기 가능", "미가입 블로그는 읽기 전용", "작성자 ID는 세션 사용자 기준으로 강제"]),
        ("슈퍼관리자 정책 기능", ["관리자 승인/권한 부여", "게시판 생성 쿼터 제한", "이미지 개수/크기 제한", "관리 차단/프로파일 제한"]),
        ("데이터 모델 설계 포인트", ["blog_user: 역할/승인/권한/제한/프로파일", "blog: 게시판/이미지 제한 값", "기본값(default) 설계가 운영 안정성에 중요"]),
        ("마이그레이션 전략", ["신규 구축: 01_create_tables.sql", "기존 업그레이드: 03_migrate_to_super_admin_model.sql", "role 정규화: ADMIN->SUPER_ADMIN, USER->BLOG_MEMBER"]),
        ("백엔드 아키텍처 연결", ["Controller -> Service -> Mapper -> DB", "Interceptor로 /admin/** 보호", "권한 체크는 서버 계층에서 일관 적용"]),
        ("화면 설계와 UX", ["관리자 대시보드에서 정책값 편집", "읽기 전용 사용자에게 안내 메시지 제공", "승인 대기 상태를 명확히 표시"]),
        ("테스트 전략", ["기능 테스트: 로그인/승인/권한 변경", "보안 테스트: 위변조/우회 시도 차단", "회귀 테스트: 기존 기능 영향 확인"]),
        ("수용 기준(AC) 예시", ["승인 전 BLOG_ADMIN은 대기 화면 이동", "SUPER_ADMIN만 /admin/** 접근 가능", "가입 블로그 외 글쓰기 차단", "정책 저장 후 DB 반영"]),
        ("워터폴 vs 애자일", ["워터폴: 문서 중심, 변경 대응 느림", "애자일: 반복 중심, 피드백 빠름", "수업 프로젝트 권장: 애자일 + 최소 PRD"]),
        ("스프린트 운영 예시", ["스프린트1: 인증/역할/마이그레이션", "스프린트2: 관리자 정책 UI/권한 검증", "리뷰/회고로 개선점 도출"]),
        ("학생 실습 확장 과제", ["댓글/대댓글 권한 모델 추가", "게시글 비밀번호 열람 정책 구현", "감사 로그 테이블 설계"]),
        ("결론 및 Q&A", ["PRD는 개발 전체의 기준선", "권한 기능의 최종 책임은 서버", "좋은 요구사항이 좋은 구현을 만든다"]),
    ]

    page_no = 2
    for title, bullets in slides:
        add_content_slide(prs, page_no, title, bullets)
        page_no += 1

    prs.save(OUT_PATH)
    print("Saved:", OUT_PATH)


if __name__ == "__main__":
    main()
