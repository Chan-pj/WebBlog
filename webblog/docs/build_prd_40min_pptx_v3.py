# -*- coding: utf-8 -*-
"""Generate styled 40-minute PRD slides v3 (cover metadata + optional logo/screenshots)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT_PATH = r"d:\eGov4.3.1\workspace-egov\webblog\docs\수업발표용_40분_20페이지_블로그PRD_v3.pptx"
SCREENSHOT_DIR = Path(r"d:\eGov4.3.1\workspace-egov\webblog\docs\screenshots")
LOGO_CANDIDATES = [
    Path(r"d:\eGov4.3.1\workspace-egov\webblog\docs\school-logo.png"),
    Path(r"d:\eGov4.3.1\workspace-egov\webblog\docs\school-logo.jpg"),
]

SCHOOL_NAME = "한국폴리텍대학"
COURSE_NAME = "블로그 실습"
PROFESSOR_NAME = "박영길"
SEMESTER = "2026년1학기"

COLOR_PRIMARY = RGBColor(28, 63, 170)
COLOR_ACCENT = RGBColor(34, 139, 230)
COLOR_TEXT = RGBColor(33, 37, 41)
COLOR_LIGHT_BG = RGBColor(245, 248, 255)


def add_footer(slide, page_text):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_LIGHT_BG
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(10.1), Inches(7.18), Inches(2.8), Inches(0.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = page_text
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY


def add_cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_LIGHT_BG
    bg.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.9), Inches(1.0), Inches(9.8), Inches(1.6))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "PRD 기반 블로그 플랫폼 설계와 구현"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    sub = slide.shapes.add_textbox(Inches(0.9), Inches(2.5), Inches(9.8), Inches(1.2))
    st = sub.text_frame
    st.clear()
    sp = st.paragraphs[0]
    sp.text = f"{SCHOOL_NAME} | {COURSE_NAME}"
    sp.font.size = Pt(20)
    sp.font.color.rgb = COLOR_TEXT
    sp2 = st.add_paragraph()
    sp2.text = f"{PROFESSOR_NAME} | {SEMESTER}"
    sp2.font.size = Pt(17)
    sp2.font.color.rgb = COLOR_TEXT

    logo_path = next((p for p in LOGO_CANDIDATES if p.exists()), None)
    if logo_path:
        slide.shapes.add_picture(str(logo_path), Inches(10.4), Inches(0.9), width=Inches(2.3))
    else:
        ph = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.4), Inches(0.9), Inches(2.3), Inches(1.2))
        ph.fill.solid()
        ph.fill.fore_color.rgb = RGBColor(225, 235, 255)
        ph.line.fill.background()
        ptf = ph.text_frame
        ptf.text = "학교 로고\n(선택)"
        ptf.paragraphs[0].font.size = Pt(12)
        ptf.paragraphs[0].font.color.rgb = COLOR_PRIMARY

    add_footer(slide, "1/20")


def add_content_slide(prs, page_text, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.16), Inches(12), Inches(0.6))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(255, 255, 255)

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.22), Inches(0.12), Inches(5.6))
    accent.fill.solid()
    accent.fill.fore_color.rgb = COLOR_ACCENT
    accent.line.fill.background()

    body = slide.shapes.add_textbox(Inches(0.9), Inches(1.25), Inches(11.9), Inches(5.6))
    tf = body.text_frame
    tf.clear()
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {line}"
        p.font.size = Pt(22)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(12)

    add_footer(slide, page_text)


def add_screenshot_slide(prs, page_text, image_path, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.16), Inches(12), Inches(0.6))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(24)
    tp.font.bold = True
    tp.font.color.rgb = RGBColor(255, 255, 255)

    slide.shapes.add_picture(str(image_path), Inches(0.8), Inches(1.2), width=Inches(11.8), height=Inches(5.7))
    add_footer(slide, page_text)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_cover_slide(prs)

    contents = [
        ("2/20", "왜 PRD가 필요한가", ["요구사항 누락과 해석 차이를 줄인다", "팀의 공통 언어를 만든다", "테스트 기준(수용 기준)을 명확히 한다"]),
        ("3/20", "PRD의 핵심 구성요소", ["문제 정의와 목표", "범위(In Scope / Out of Scope)", "사용자 역할/시나리오", "기능·비기능 요구사항", "수용 기준과 성공 지표"]),
        ("4/20", "프로젝트 배경과 문제", ["권한 분리가 약하면 운영 안정성이 떨어진다", "승인 절차가 없으면 관리자 권한 오남용 가능", "업로드 정책이 없으면 저장소/비용 통제가 어렵다"]),
        ("5/20", "역할 모델(RBAC) 설계", ["SUPER_ADMIN", "BLOG_ADMIN", "BLOG_MEMBER", "승인 상태: is_approved"]),
        ("6/20", "인증/가입 요구사항", ["이메일 로그인", "카카오 OAuth 로그인", "가입 유형 선택(BLOG_ADMIN / BLOG_MEMBER)", "BLOG_ADMIN은 승인 전 대기"]),
        ("7/20", "로그인 후 라우팅 정책", ["SUPER_ADMIN -> 관리자 대시보드", "승인 전 BLOG_ADMIN -> 대기 화면", "BLOG_MEMBER -> 소속 블로그"]),
        ("8/20", "블로그멤버 글쓰기 정책", ["소유자 또는 ACTIVE 멤버만 글쓰기 가능", "미가입 블로그는 읽기 전용", "작성자 ID는 세션 사용자 기준으로 강제"]),
        ("9/20", "슈퍼관리자 정책 기능", ["관리자 승인/권한 부여", "게시판 생성 쿼터 제한", "이미지 개수/크기 제한", "관리 차단/프로파일 제한"]),
        ("10/20", "데이터 모델 설계 포인트", ["blog_user: 역할/승인/권한/제한/프로파일", "blog: 게시판/이미지 제한 값", "기본값(default) 설계가 운영 안정성에 중요"]),
        ("11/20", "마이그레이션 전략", ["신규 구축: 01_create_tables.sql", "기존 업그레이드: 03_migrate_to_super_admin_model.sql", "role 정규화: ADMIN->SUPER_ADMIN, USER->BLOG_MEMBER"]),
        ("12/20", "백엔드 아키텍처 연결", ["Controller -> Service -> Mapper -> DB", "Interceptor로 /admin/** 보호", "권한 체크는 서버 계층에서 일관 적용"]),
        ("13/20", "화면 설계와 UX", ["관리자 대시보드에서 정책값 편집", "읽기 전용 사용자에게 안내 메시지 제공", "승인 대기 상태를 명확히 표시"]),
        ("14/20", "테스트 전략", ["기능 테스트: 로그인/승인/권한 변경", "보안 테스트: 위변조/우회 시도 차단", "회귀 테스트: 기존 기능 영향 확인"]),
        ("15/20", "수용 기준(AC) 예시", ["승인 전 BLOG_ADMIN은 대기 화면 이동", "SUPER_ADMIN만 /admin/** 접근 가능", "가입 블로그 외 글쓰기 차단", "정책 저장 후 DB 반영"]),
    ]
    for c in contents:
        add_content_slide(prs, *c)

    # 16~18: optional screenshot slides
    images = []
    if SCREENSHOT_DIR.exists():
        for ext in ("*.png", "*.jpg", "*.jpeg", "*.webp"):
            images.extend(sorted(SCREENSHOT_DIR.glob(ext)))
    images = images[:3]
    screenshot_titles = ["예시 화면 1", "예시 화면 2", "예시 화면 3"]
    for idx in range(3):
        page = f"{16 + idx}/20"
        if idx < len(images):
            add_screenshot_slide(prs, page, images[idx], screenshot_titles[idx])
        else:
            add_content_slide(
                prs,
                page,
                screenshot_titles[idx],
                ["`docs/screenshots`에 이미지를 넣으면 자동 삽입됩니다.", "추천 파일명: 01-home.png, 02-login.png, 03-admin-dashboard.png"],
            )

    add_content_slide(prs, "19/20", "학생 실습 확장 과제", ["댓글/대댓글 권한 모델 추가", "게시글 비밀번호 열람 정책 구현", "감사 로그 테이블 설계"])
    add_content_slide(prs, "20/20", "결론 및 Q&A", ["PRD는 개발 전체의 기준선", "권한 기능의 최종 책임은 서버", "좋은 요구사항이 좋은 구현을 만든다"])

    prs.save(OUT_PATH)
    print("Saved:", OUT_PATH)


if __name__ == "__main__":
    main()
