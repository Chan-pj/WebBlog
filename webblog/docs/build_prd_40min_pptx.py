# -*- coding: utf-8 -*-
"""Generate 40-minute / 20-page PRD lecture slides."""
from pptx import Presentation
from pptx.util import Inches, Pt

OUT_PATH = r"d:\eGov4.3.1\workspace-egov\webblog\docs\수업발표용_40분_20페이지_블로그PRD.pptx"


def add_title_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def add_bullet_slide(prs, title, bullets, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, line in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(21)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "PRD 기반 블로그 플랫폼 설계와 구현",
        "수업 발표 자료 (40분 · 20페이지)",
    )

    slides = [
        (
            "2. 왜 PRD가 필요한가",
            [
                "요구사항 누락과 해석 차이를 줄인다",
                "팀의 공통 언어를 만든다",
                "테스트 기준(수용 기준)을 명확히 한다",
            ],
            "PRD는 개발 속도를 늦추는 문서가 아니라, 재작업 비용을 줄이는 문서임을 강조합니다.",
        ),
        (
            "3. PRD의 핵심 구성요소",
            [
                "문제 정의와 목표",
                "범위(In Scope / Out of Scope)",
                "사용자 역할/시나리오",
                "기능·비기능 요구사항",
                "수용 기준과 성공 지표",
            ],
            "학생들에게 PRD는 설계서가 아니라 제품 방향 문서라는 점을 설명합니다.",
        ),
        (
            "4. 프로젝트 배경과 문제",
            [
                "권한 분리가 약하면 운영 안정성이 떨어진다",
                "승인 절차가 없으면 관리자 권한 오남용 가능",
                "업로드 정책이 없으면 저장소/비용 통제가 어렵다",
            ],
            "문제 정의가 선행되어야 요구사항 우선순위를 세울 수 있음을 설명합니다.",
        ),
        (
            "5. 이해관계자와 사용자",
            [
                "슈퍼관리자: 승인/정책 통제",
                "블로그관리자: 블로그 운영",
                "블로그멤버: 열람·작성",
                "강사/조교: 학습 성과 검증",
            ],
            "역할별 기대가 다르므로 요구사항도 역할별로 작성해야 합니다.",
        ),
        (
            "6. 역할 모델(RBAC) 설계",
            [
                "SUPER_ADMIN",
                "BLOG_ADMIN",
                "BLOG_MEMBER",
                "승인 상태: is_approved",
            ],
            "역할(role)과 상태(state)를 함께 봐야 실제 권한을 정의할 수 있습니다.",
        ),
        (
            "7. 인증/가입 요구사항",
            [
                "이메일 로그인",
                "카카오 OAuth 로그인",
                "가입 유형 선택(BLOG_ADMIN / BLOG_MEMBER)",
                "BLOG_ADMIN은 승인 전 대기",
            ],
            "같은 인증 기능이라도 가입 이후 상태 전이가 제품 품질을 좌우합니다.",
        ),
        (
            "8. 로그인 후 라우팅 정책",
            [
                "SUPER_ADMIN -> 관리자 대시보드",
                "승인 전 BLOG_ADMIN -> 대기 화면",
                "BLOG_MEMBER -> 소속 블로그",
            ],
            "로그인 성공 이후의 경로 정책을 요구사항으로 명시해야 혼선이 줄어듭니다.",
        ),
        (
            "9. 블로그멤버 글쓰기 정책",
            [
                "소유자 또는 ACTIVE 멤버만 글쓰기 가능",
                "미가입 블로그는 읽기 전용",
                "작성자 ID는 세션 사용자 기준으로 강제",
            ],
            "UI 제어만으로는 보안이 안 되므로 서버 검증이 필수임을 강조합니다.",
        ),
        (
            "10. 슈퍼관리자 정책 기능",
            [
                "관리자 승인/권한 부여",
                "게시판 생성 쿼터 제한",
                "이미지 개수/크기 제한",
                "관리 차단/프로파일 제한",
            ],
            "정책을 데이터 컬럼과 화면 설정으로 연결하면 운영 자동화가 가능합니다.",
        ),
        (
            "11. 데이터 모델 설계 포인트",
            [
                "blog_user: 역할/승인/권한/제한/프로파일",
                "blog: 게시판/이미지 제한 값",
                "기본값(default) 설계가 운영 안정성에 중요",
            ],
            "정책형 서비스에서는 DB 스키마 자체가 운영 정책이 됩니다.",
        ),
        (
            "12. 마이그레이션 전략",
            [
                "신규 구축: 01_create_tables.sql",
                "기존 업그레이드: 03_migrate_to_super_admin_model.sql",
                "role 정규화: ADMIN->SUPER_ADMIN, USER->BLOG_MEMBER",
            ],
            "현업에서는 신규 개발보다 무중단/안전 마이그레이션 역량이 중요합니다.",
        ),
        (
            "13. 백엔드 아키텍처 연결",
            [
                "Controller -> Service -> Mapper -> DB",
                "Interceptor로 /admin/** 보호",
                "권한 체크는 서버 계층에서 일관 적용",
            ],
            "계층 분리가 되어야 정책 변경 시 수정 범위를 최소화할 수 있습니다.",
        ),
        (
            "14. 화면 설계와 UX",
            [
                "관리자 대시보드에서 정책값 편집",
                "읽기 전용 사용자에게 안내 메시지 제공",
                "승인 대기 상태를 명확히 표시",
            ],
            "권한 차단 시 사용자에게 이유를 설명하면 불필요한 문의를 줄일 수 있습니다.",
        ),
        (
            "15. 테스트 전략",
            [
                "기능 테스트: 로그인/승인/권한 변경",
                "보안 테스트: 위변조/우회 시도 차단",
                "회귀 테스트: 기존 기능 영향 확인",
            ],
            "권한 기능은 정상 케이스보다 비정상 케이스 테스트가 더 중요합니다.",
        ),
        (
            "16. 수용 기준(AC) 예시",
            [
                "승인 전 BLOG_ADMIN은 대기 화면 이동",
                "SUPER_ADMIN만 /admin/** 접근 가능",
                "가입 블로그 외 글쓰기 차단",
                "정책 저장 후 DB 반영",
            ],
            "AC는 '언제 완료로 볼 것인가'를 정의하는 합의 장치입니다.",
        ),
        (
            "17. 워터폴 vs 애자일",
            [
                "워터폴: 문서 중심, 변경 대응 느림",
                "애자일: 반복 중심, 피드백 빠름",
                "수업 프로젝트 권장: 애자일 + 최소 PRD",
            ],
            "방법론은 정답이 아니라 상황에 맞춘 선택임을 설명합니다.",
        ),
        (
            "18. 스프린트 운영 예시",
            [
                "스프린트1: 인증/역할/마이그레이션",
                "스프린트2: 관리자 정책 UI/권한 검증",
                "리뷰/회고로 개선점 도출",
            ],
            "작은 완료 단위를 빠르게 반복하면 팀 학습 속도가 빨라집니다.",
        ),
        (
            "19. 학생 실습 확장 과제",
            [
                "댓글/대댓글 권한 모델 추가",
                "게시글 비밀번호 열람 정책 구현",
                "감사 로그 테이블 설계",
            ],
            "오늘 배운 PRD 템플릿으로 기능 하나를 직접 정의해보게 유도합니다.",
        ),
        (
            "20. 결론 및 Q&A",
            [
                "PRD는 개발 전체의 기준선",
                "권한 기능의 최종 책임은 서버",
                "좋은 요구사항이 좋은 구현을 만든다",
            ],
            "질의응답 시 '요구사항-설계-구현-테스트' 연결 관점으로 답변하도록 유도합니다.",
        ),
    ]

    for title, bullets, notes in slides:
        add_bullet_slide(prs, title, bullets, notes)

    prs.save(OUT_PATH)
    print("Saved:", OUT_PATH)


if __name__ == "__main__":
    main()
